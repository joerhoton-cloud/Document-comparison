"""Generate sample original/revised contracts (DOCX, PDF, TXT, scanned PDF) for trying the app.

Requires: pip install python-docx reportlab pillow
Run from this folder: python3 generate.py
"""
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ORIGINAL = [
    ("h", "Media Monitoring Services Agreement"),
    ("p", 'This Media Monitoring Services Agreement (the "Agreement") is entered into as of January 1, 2026 by and between Provider Inc. ("Provider") and Example Corp. ("Customer").'),
    ("h", "1. Services"),
    ("p", "Provider shall deliver daily media monitoring reports covering print, broadcast and online sources in the United States and Canada."),
    ("p", "Reports will be delivered by email no later than 8:00 a.m. Eastern Time each business day."),
    ("h", "2. Fees and Payment"),
    ("p", "Customer shall pay an annual subscription fee of $48,000, payable quarterly in advance. Invoices are due within thirty (30) days of receipt."),
    ("p", "Late payments shall accrue interest at the rate of 1.5% per month."),
    ("h", "3. Term and Termination"),
    ("p", "The initial term of this Agreement is twelve (12) months and shall automatically renew for successive one-year terms unless either party gives sixty (60) days written notice."),
    ("p", "Either party may terminate this Agreement for material breach that remains uncured thirty (30) days after written notice."),
    ("h", "4. Confidentiality"),
    ("p", "Each party shall protect the other party's Confidential Information using at least the same degree of care it uses to protect its own confidential information, but no less than reasonable care."),
    ("h", "5. Limitation of Liability"),
    ("p", "In no event shall either party's aggregate liability exceed the fees paid by Customer in the twelve (12) months preceding the claim."),
    ("h", "6. Governing Law"),
    ("p", "This Agreement shall be governed by the laws of the State of Illinois."),
]

REVISED = [
    ("h", "Media Monitoring Services Agreement"),
    ("p", 'This Media Monitoring Services Agreement (the "Agreement") is entered into as of February 1, 2026 by and between Provider Inc. ("Provider") and Example Corp. ("Customer").'),
    ("h", "1. Services"),
    ("p", "Provider shall deliver daily media monitoring reports covering print, broadcast, podcast and online sources in the United States, Canada and Mexico."),
    ("p", "Reports will be delivered through the online platform no later than 7:00 a.m. Eastern Time each business day."),
    ("p", "Provider shall also deliver a monthly executive summary with share-of-voice analysis."),
    ("h", "2. Fees and Payment"),
    ("p", "Customer shall pay an annual subscription fee of $52,500, payable annually in advance. Invoices are due within forty-five (45) days of receipt."),
    ("h", "3. Term and Termination"),
    ("p", "The initial term of this Agreement is twenty-four (24) months and shall automatically renew for successive one-year terms unless either party gives ninety (90) days written notice."),
    ("p", "Either party may terminate this Agreement for material breach that remains uncured thirty (30) days after written notice."),
    ("h", "4. Confidentiality"),
    ("p", "Each party shall protect the other party's Confidential Information using at least the same degree of care it uses to protect its own confidential information, but no less than reasonable care."),
    ("h", "5. Limitation of Liability"),
    ("p", "In no event shall either party's aggregate liability exceed the fees paid by Customer in the twelve (12) months preceding the claim."),
    ("h", "6. Governing Law"),
    ("p", "This Agreement shall be governed by the laws of the State of Delaware."),
]


def make_docx(items, path):
    d = Document()
    for kind, text in items:
        if kind == "h":
            d.add_heading(text, level=2 if text[0].isdigit() else 1)
        else:
            d.add_paragraph(text)
    d.save(path)


def make_pdf(items, path):
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=letter)
    story = []
    for kind, text in items:
        story.append(Paragraph(text.replace("&", "&amp;"), styles["Heading2" if kind == "h" else "BodyText"]))
        story.append(Spacer(1, 6))
    doc.build(story)


def make_scanned_pdf(items, path, dpi=200):
    """An image-only PDF (no text layer), like a photocopier scan: slight tilt and speckle."""
    import random
    import textwrap
    from PIL import Image, ImageDraw, ImageFont

    random.seed(7)
    w, h = int(8.5 * dpi), int(11 * dpi)
    margin = int(1 * dpi)
    body = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", int(dpi * 0.15))
    head = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf", int(dpi * 0.19))
    pages, page, y = [], None, h
    def new_page():
        nonlocal page, y
        page = Image.new("L", (w, h), 250)
        pages.append(page)
        y = margin
    for kind, text in items:
        font = head if kind == "h" else body
        lines = textwrap.wrap(text, 70 if kind == "p" else 60)
        line_h = int(font.size * 1.45)
        if y + line_h * len(lines) > h - margin:
            new_page()
        if page is None:
            new_page()
        draw = ImageDraw.Draw(page)
        for line in lines:
            draw.text((margin, y), line, font=font, fill=25)
            y += line_h
        y += int(font.size * 0.9)
    out = []
    for pg in pages:
        px = pg.load()
        for _ in range(1500):  # scanner speckle
            px[random.randrange(w), random.randrange(h)] = random.randrange(80, 200)
        out.append(pg.rotate(0.4, fillcolor=250, resample=Image.BICUBIC).convert("RGB"))
    if path.endswith(".png"):
        out[0].save(path)  # first page only, like a phone photo of one page
    else:
        out[0].save(path, save_all=True, append_images=out[1:], resolution=dpi)


if __name__ == "__main__":
    make_docx(ORIGINAL, "contract-original.docx")
    make_docx(REVISED, "contract-revised.docx")
    make_pdf(ORIGINAL, "contract-original.pdf")
    make_pdf(REVISED, "contract-revised.pdf")
    make_scanned_pdf(ORIGINAL, "contract-original-scanned.pdf")
    make_scanned_pdf(REVISED, "contract-revised-page1.png")
    with open("contract-original.txt", "w") as f:
        f.write("\n\n".join(t for _, t in ORIGINAL) + "\n")
